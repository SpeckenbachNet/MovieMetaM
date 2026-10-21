#!/usr/bin/env python3
# Copyright (C) 2026 Michael Speckenbach – GPL v3 or later (see LICENSE)
"""Writes the age rating systems of MovieMetaM from the working table.

    python3 tools/ratings_from_xlsx.py ~/kDrive/Development/QT/age_ratings.xlsx

Reads the sheet "Ratings" (one row per step) and writes one JSON file per country to
ressources/ratings/<code>.json plus the file list in Ratings.qrc and countrynames.cpp (the country
names for lupdate: one translatable text each, context "Country"). Needs openpyxl
(pip install --user openpyxl). Run it again after every change of the table.
"""
import json
import os
import sys
from collections import OrderedDict

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "ressources", "ratings")
QRC = os.path.join(ROOT, "Ratings.qrc")
NAMES_CPP = os.path.join(ROOT, "countrynames.cpp")


def text(value):
    return "" if value is None else str(value).strip()


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    sheet = load_workbook(sys.argv[1], read_only=True, data_only=True)["Ratings"]
    rows = sheet.iter_rows(values_only=True)
    head = [text(h) for h in next(rows)]
    col = {name: head.index(name) for name in head}

    countries = OrderedDict()
    for row in rows:
        code = text(row[col["Code"]])
        if not code:
            continue
        country = countries.setdefault(code, {
            "country": {"en": text(row[col["Country (en)"]]), "de": text(row[col["Country (de)"]])},
            "system": text(row[col["System"]]),
            "id": text(row[col["Prefix (id)"]]),
            "ratings": [],
        })
        description = text(row[col["Description (tooltip)"]])
        country["ratings"].append({
            "key": text(row[col["Badge text"]]),                      # text in the badge
            "code": text(row[col["Code to write (media servers)"]]),  # written into the file
            "score": int(row[col["Score"]]),
            "label": description or text(row[col["Apple / logo code"]]),   # text in the combo box
            "description": description,                               # tooltip
            "color": text(row[col["Color"]]).upper(),
            "shape": text(row[col["Shape"]]),
            "style": text(row[col["Style"]]) or "filled",
        })

    os.makedirs(OUT_DIR, exist_ok=True)
    for name in os.listdir(OUT_DIR):   # countries removed from the table disappear
        if name.endswith(".json"):
            os.remove(os.path.join(OUT_DIR, name))
    for code, country in countries.items():
        country["ratings"].sort(key=lambda r: r["score"])
        with open(os.path.join(OUT_DIR, f"{code}.json"), "w", encoding="utf-8") as f:
            json.dump(country, f, ensure_ascii=False, indent=4)
            f.write("\n")

    with open(QRC, "w", encoding="utf-8") as f:
        f.write("<RCC>\n")
        f.write("    <!-- Age rating systems, one file per country: <country code>.json (see agerating.h).\n")
        f.write("         Written by tools/ratings_from_xlsx.py – do not edit by hand. -->\n")
        f.write('    <qresource prefix="/ratings">\n')
        for code in sorted(countries):
            f.write(f'        <file alias="{code}.json">ressources/ratings/{code}.json</file>\n')
        f.write("    </qresource>\n</RCC>\n")

    # The names of the countries as texts for lupdate: translated in the .ts files (context
    # "Country"), the English name is the key (see agerating.cpp)
    names = []
    for country in countries.values():
        name = country["country"]["en"]
        if name and name not in names:
            names.append(name)
    with open(NAMES_CPP, "w", encoding="utf-8") as f:
        f.write("// Written by tools/ratings_from_xlsx.py – do not edit by hand.\n")
        f.write("// The names of the countries of the age rating systems, only here for lupdate: they are\n")
        f.write("// translated in the .ts files (context \"Country\") and looked up in agerating.cpp.\n")
        f.write("#include <QtGlobal>\n\n")
        f.write("[[maybe_unused]] static const char *const kCountryNames[] = {\n")
        for name in names:
            f.write(f'    QT_TRANSLATE_NOOP("Country", "{name}"),\n')
        f.write("};\n")

    print(f"{len(countries)} countries, {sum(len(c['ratings']) for c in countries.values())} steps written")


if __name__ == "__main__":
    main()
